"""Generates the historical dossiers workbook for the import page (Reclamations + Causes sheets).

Deterministic (fixed seed). Vocabulary follows the existing catalogue (defects, lines, clients,
shifts, lot format) so the history reads like the rest of the base. Dates run from 2023-01 to
2025-06, i.e. before the claims already in the database.

Usage: python generate_history_workbook.py <output.xlsx> [count]
"""
import random
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = sys.argv[1] if len(sys.argv) > 1 else "historique-dossiers.xlsx"
COUNT = int(sys.argv[2]) if len(sys.argv) > 2 else 180
rng = random.Random(20230117)

START, END = date(2023, 1, 9), date(2025, 6, 27)

CLIENTS = ["BMW-DEB-4521", "BMW-REG-3301", "VW-AUD-1180", "VW-WOL-2044", "STELLANTIS-MUL", "PSA-SOCH-220",
           "RENAULT-FLINS", "RENAULT-SAN", "DAIMLER-STUT", "FORD-VALENCIA", "TOYOTA-VAL", "VALEO-TUNIS",
           "LEONI-SOUSSE", "YAZAKI-BIZERTE", "APTIV-MOROCCO", "SUMITOMO-TN"]
SHIFTS = ["Matin", "Après-midi", "Nuit", "Équipe A", "Équipe B"]
ZONES = ["A12", "B07", "C03", "ancrage", "pin 4", "pin 12", "branche gauche", "branche droite", "feuillard"]
SIDES = ["gauche", "droit", "central"]

# defect -> (weight, [(line, weight)], preferred shifts, title templates, description template)
DEFECTS = {
    "Soudure défectueuse": (22, [("WR-04", 55), ("WR-05", 25), ("WR-02", 12), ("WR-01", 8)], ["Nuit", "Équipe B"],
        ["Soudure {zone} ancrage {side} — robot {line}", "Soudure froide connecteur {zone} — {line}", "Récurrence soudure ancrage — {line}"],
        "Point de soudure {zone} non conforme sur faisceau client {client}. Aspect poreux / manque de pénétration. Lot {lot}. Ligne {line}, équipe {shift}."),
    "Isolation endommagée": (18, [("EXT-01", 45), ("CUT-02", 30), ("EXT-02", 25)], ["Matin", "Équipe A"],
        ["Isolation abrasée sur {zone} — extrusion {line}", "Coupure partielle isolation — {line}"],
        "Gaine d'isolation endommagée (abrasion / pincement) sur tronçon {zone}. Risque court-circuit. Détection en fin de ligne {line}. Lot {lot}, client {client}."),
    "Dimension hors tolérance": (14, [("CUT-01", 60), ("CUT-02", 40)], ["Après-midi", "Matin"],
        ["Longueur hors tolérance ±{tol} mm — {line}", "Cote branche {zone} hors tolérance — {line}"],
        "Mesure hors cote sur longueur de branche {zone}. Tolérance dépassée. Ligne {line}, lot {lot}, client {client}."),
    "Coupure fil": (12, [("ASM-03", 50), ("ASM-01", 30), ("CUT-01", 20)], ["Nuit", "Après-midi"],
        ["Fil sectionné au sertissage — {line}", "Conducteur coupé zone {zone} — {line}"],
        "Conducteur sectionné au niveau de la zone de sertissage {zone}. Client {client}, lot {lot}."),
    "Manque de sertissage": (11, [("ASM-01", 60), ("ASM-03", 40)], ["Équipe B", "Nuit"],
        ["Sertissage incomplet terminal {zone}", "Hauteur de sertissage hors fenêtre — {line}"],
        "Hauteur / largeur de sertissage hors fenêtre. Terminal {zone} mal formé. Ligne {line}, lot {lot}, client {client}."),
    "Débordement étain": (10, [("WR-02", 55), ("WR-01", 25), ("WR-04", 20)], ["Matin", "Après-midi"],
        ["Débordement étain soudure {zone}", "Excès d'étain plots {zone} — {line}"],
        "Excès d'étain avec risque de pontage entre plots. Ligne {line}, lot {lot}."),
    "Connecteur desserré": (9, [("ASM-03", 65), ("ASM-01", 35)], ["Équipe A", "Matin"],
        ["Connecteur {zone} mal verrouillé — {line}"],
        "CPA / levier de verrouillage non enclenché. Risque de déconnexion en montage client. Lot {lot}, client {client}."),
    "Marquage illisible": (8, [("TEST-01", 70), ("EXT-02", 30)], ["Après-midi"],
        ["Marquage inkjet illisible — lot {lot}"],
        "Code lot / référence illisible après marquage. Ligne {line}."),
    "Couleur incorrecte": (7, [("EXT-02", 80), ("EXT-01", 20)], ["Matin"],
        ["Couleur gaine incorrecte — {zone}"],
        "Couleur de gaine non conforme au plan (attendu vs reçu). Client {client}, lot {lot}."),
    "Contamination surface": (7, [("WR-01", 50), ("WR-03", 30), ("WR-06", 20)], ["Nuit", "Après-midi"],
        ["Contamination huile / particules — {line}"],
        "Traces de contamination sur zone de soudure {zone}. Lot {lot}."),
    "Non-conformité emballage": (6, [("ASM-01", 60), ("TEST-01", 40)], ["Après-midi", "Équipe B"],
        ["Emballage non conforme expédition — {client}"],
        "Conditionnement hors standard client (séparateurs / étiquette manquante). Lot {lot} bloqué quai expédition."),
    "Étiquette manquante": (5, [("TEST-01", 80), ("ASM-01", 20)], ["Matin", "Équipe A"],
        ["Étiquette absente lot {lot}"],
        "Étiquette identification produit manquante sur carton / bobine. Client {client}."),
    "Polarité inversée": (4, [("ASM-03", 75), ("ASM-01", 25)], ["Nuit"],
        ["Polarité inversée circuit {zone}"],
        "Inversion de polarité détectée au test électrique. Ligne {line}, lot {lot}."),
    "Longueur câble incorrecte": (4, [("CUT-02", 70), ("CUT-01", 30)], ["Équipe B", "Après-midi"],
        ["Longueur câble hors cote — {line}"],
        "Longueur mesurée hors tolérance sur référence client {client}. Lot {lot}."),
}

# defect -> list of (category 6M, cause text, probability of appearing, can be the root cause)
CAUSES = {
    "Soudure défectueuse": [
        ("Machine", "Usure des électrodes de soudure non détectée à temps", 0.8, True),
        ("Machine", "Déréglage de la pression de la pince", 0.45, True),
        ("Méthode", "Paramètres de soudure (pression / temps) non standardisés entre équipes", 0.5, False),
        ("Matière", "Lot de feuillard avec épaisseur en limite de tolérance", 0.25, False),
        ("Main-d'œuvre", "Changement d'électrodes réalisé sans la procédure de contrôle", 0.3, False),
        ("Mesure", "Contrôle de résistance en ligne peu sensible à la soudure froide", 0.3, False)],
    "Isolation endommagée": [
        ("Machine", "Guide d'extrusion usé créant un point de frottement", 0.65, True),
        ("Machine", "Lame de coupe émoussée", 0.45, True),
        ("Méthode", "Fréquence de remplacement des lames non définie", 0.4, False),
        ("Milieu", "Poussières abrasives sur le chemin de câble", 0.25, False),
        ("Main-d'œuvre", "Manipulation des bobines sans protection", 0.3, False)],
    "Dimension hors tolérance": [
        ("Machine", "Réglage de la table de coupe dérivé après changement de référence", 0.7, True),
        ("Mesure", "Gabarit de mesure non étalonné", 0.45, True),
        ("Méthode", "Première pièce non validée après réglage", 0.45, False),
        ("Main-d'œuvre", "Saisie manuelle de la longueur erronée", 0.3, False)],
    "Coupure fil": [
        ("Machine", "Hauteur de sertissage trop serrée sur l'applicateur", 0.6, True),
        ("Matière", "Conducteur fragilisé sur une bobine du fournisseur", 0.35, True),
        ("Méthode", "Contrôle de traction du sertissage trop espacé", 0.4, False),
        ("Main-d'œuvre", "Réglage de l'applicateur par un opérateur non habilité", 0.3, False)],
    "Manque de sertissage": [
        ("Machine", "Réglage hauteur de sertissage hors fenêtre après changement d'applicateur", 0.75, True),
        ("Méthode", "Fiche de réglage applicateur incomplète", 0.45, False),
        ("Mesure", "Contrôle crimp force non réalisé en début de série", 0.4, False),
        ("Matière", "Terminal hors spécification dimensionnelle", 0.25, True)],
    "Débordement étain": [
        ("Machine", "Température du four au-dessus de la fenêtre", 0.6, True),
        ("Méthode", "Quantité d'étain non limitée par le mode opératoire", 0.4, False),
        ("Main-d'œuvre", "Retouche manuelle sans gabarit", 0.35, False)],
    "Connecteur desserré": [
        ("Méthode", "Contrôle du verrouillage CPA absent du poste", 0.6, True),
        ("Main-d'œuvre", "Cadence élevée en fin d'équipe", 0.4, False),
        ("Matière", "Connecteur avec levier déformé", 0.3, True)],
    "Marquage illisible": [
        ("Machine", "Tête d'impression encrassée", 0.7, True),
        ("Matière", "Encre hors date de validité", 0.3, True),
        ("Méthode", "Nettoyage de la tête non planifié", 0.4, False)],
    "Couleur incorrecte": [
        ("Matière", "Mauvais lot de mélange maître chargé en extrusion", 0.55, True),
        ("Méthode", "Contrôle couleur au changement de bobine non formalisé", 0.5, False),
        ("Main-d'œuvre", "Erreur de chargement trémie", 0.35, False)],
    "Contamination surface": [
        ("Milieu", "Zone de soudure exposée aux projections d'huile", 0.5, True),
        ("Méthode", "Nettoyage du poste non tracé", 0.45, False),
        ("Main-d'œuvre", "Gants non remplacés entre les références", 0.4, True)],
    "Non-conformité emballage": [
        ("Méthode", "Spécification emballage client non reprise dans l'instruction de poste", 0.6, True),
        ("Main-d'œuvre", "Séparateurs omis en fin de conditionnement", 0.4, False),
        ("Mesure", "Pas de contrôle emballage avant expédition", 0.35, False)],
    "Étiquette manquante": [
        ("Méthode", "Étape d'étiquetage sans contrôle de présence", 0.6, True),
        ("Machine", "Imprimante d'étiquettes en erreur sans alarme", 0.3, True),
        ("Main-d'œuvre", "Rouleau d'étiquettes épuisé non remplacé", 0.35, False)],
    "Polarité inversée": [
        ("Main-d'œuvre", "Erreur d'insertion des cavités sur le plan de câblage", 0.55, True),
        ("Méthode", "Plan de câblage ambigu sur la révision en cours", 0.4, True),
        ("Mesure", "Test électrique ne couvre pas toutes les polarités", 0.35, False)],
    "Longueur câble incorrecte": [
        ("Machine", "Dérive du codeur de longueur de la table de coupe", 0.6, True),
        ("Mesure", "Contrôle longueur par échantillonnage trop espacé", 0.4, False),
        ("Méthode", "Recoupe sans nouvelle validation", 0.3, False)],
}

SEVERITY_DEFECTIVE = {"Faible": (2, 12), "Moyenne": (6, 22), "Haute": (14, 40), "Critique": (25, 90)}


def pick(weighted):
    total = sum(w for _, w in weighted)
    r = rng.uniform(0, total)
    acc = 0
    for v, w in weighted:
        acc += w
        if r <= acc:
            return v
    return weighted[-1][0]


def weekday(d):
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d


def sample_date():
    """More claims in the second half of the period (the plant grows / traces better), one WR-04 spike."""
    span = (END - START).days
    r = rng.random() ** 0.8
    return weekday(START + timedelta(days=int(r * span)))


rows, cause_rows = [], []
defects = [(d, v[0]) for d, v in DEFECTS.items()]
used_titles = set()
spike_remaining = 14  # a cluster of WR-04 solder defects in spring 2025, then corrected

i = 0
while len(rows) < COUNT:
    i += 1
    spike = spike_remaining > 0 and rng.random() < 0.12
    if spike:
        defect = "Soudure défectueuse"
        line = "WR-04"
        detected = weekday(date(2025, 3, 3) + timedelta(days=rng.randint(0, 50)))
        spike_remaining -= 1
    else:
        defect = pick(defects)
        line = pick(DEFECTS[defect][1])
        detected = sample_date()

    _, _, shifts, titles, desc = DEFECTS[defect]
    shift = rng.choice(shifts) if rng.random() < 0.7 else rng.choice(["Matin", "Après-midi", "Nuit", "Équipe A", "Équipe B"])
    severity = rng.choices(["Faible", "Moyenne", "Haute", "Critique"], weights=[22, 38, 29, 11])[0]
    client = rng.choice(CLIENTS)
    lot = f"LOT-{detected:%Y}-{detected.timetuple().tm_yday:03d}-{rng.randint(10, 99):02d}"
    vars_ = {"zone": rng.choice(ZONES), "side": rng.choice(SIDES), "line": line, "client": client,
             "lot": lot, "shift": shift, "tol": str(rng.randint(2, 9))}
    title = rng.choice(titles).format(**vars_)
    if title in used_titles and rng.random() < 0.6:
        title = f"{defect} — {line} / {vars_['zone']}"
    used_titles.add(title)
    description = desc.format(**vars_)

    base = {"Faible": 8, "Moyenne": 14, "Haute": 22, "Critique": 32}[severity]
    closed = weekday(detected + timedelta(days=int(base + rng.uniform(-4, 18))))
    if closed > date(2025, 7, 31):
        closed = weekday(date(2025, 7, 31) - timedelta(days=rng.randint(0, 20)))

    produced = rng.randint(800, 5000)
    lo, hi = SEVERITY_DEFECTIVE[severity]
    defective = min(produced, rng.randint(lo, hi))

    ref = f"HIS-{len(rows) + 1:04d}"
    rows.append([ref, title, client, defect, severity, lot, description, detected.isoformat(), closed.isoformat(),
                 line, shift, produced, defective])

    pool = CAUSES[defect]
    chosen = [c for c in pool if rng.random() < c[2]]
    if not chosen:
        chosen = [pool[0]]
    roots = [c for c in chosen if c[3]]
    root = roots[0] if roots else chosen[0]
    for cat, text, _, _ in chosen:
        cause_rows.append([ref, cat, text, "oui" if (cat, text) == (root[0], root[1]) else "non"])

CLAIM_HEADERS = ["Reference", "Titre", "Client", "Type de defaut", "Gravite", "Numero de lot", "Description",
                 "Date de detection", "Date de cloture", "Ligne de production", "Equipe",
                 "Quantite produite", "Quantite defectueuse"]
CAUSE_HEADERS = ["Reference", "Categorie 6M", "Description de la cause", "Cause racine (oui/non)"]

wb = Workbook()
wb.remove(wb.active)
for name, headers, data, widths in (
    ("Reclamations", CLAIM_HEADERS, rows, [11, 48, 18, 26, 10, 20, 70, 16, 16, 18, 12, 16, 18]),
    ("Causes", CAUSE_HEADERS, cause_rows, [11, 16, 70, 22]),
):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8ECF8")
    for r in data:
        ws.append(r)
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = w

wb.save(OUT)
print(f"{len(rows)} dossiers, {len(cause_rows)} causes -> {OUT}")
print("periode detection:", min(r[7] for r in rows), "->", max(r[7] for r in rows))
